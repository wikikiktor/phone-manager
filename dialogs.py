import tkinter as tk
from tkinter import messagebox, ttk

import openpyxl

import db


class AddEventDialog(tk.Toplevel):
    def __init__(self, parent, phone_id, on_save_callback):
        super().__init__(parent)
        self.phone_id = phone_id
        self.on_save_callback = on_save_callback

        self.title("Dodaj zdarzenie / uwagę")
        self.geometry("450x300")
        self.transient(parent)
        self.grab_set()

        self.build_ui()

    def build_ui(self):
        ttk.Label(self, text="Kategoria:").pack(anchor=tk.W, padx=15, pady=(15, 2))
        
        self.cat_combo = ttk.Combobox(
            self,
            values=[
                "Notatka / Uwaga",
                "Zmiana użytkownika",
                "Awaria / Serwis",
                "Wydanie",
                "Zwrot",
                "Wymiana karty SIM",
                "Inne",
            ],
            state="readonly",
        )
        self.cat_combo.set("Notatka / Uwaga")
        self.cat_combo.pack(fill=tk.X, padx=15)

        ttk.Label(self, text="Opis zdarzenia / uwagi:").pack(anchor=tk.W, padx=15, pady=(10, 2))
        
        self.txt = tk.Text(self, height=5, wrap=tk.WORD)
        self.txt.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)
        self.txt.focus()

        ttk.Button(self, text="Zapisz wpis", command=self.save_event).pack(
            pady=10, padx=15, anchor=tk.E
        )

    def save_event(self):
        opis = self.txt.get("1.0", tk.END).strip()

        if not opis:
            messagebox.showwarning("Puste pole", "Wpisz treść zdarzenia lub uwagi.")
            return

        db.add_event(self.phone_id, self.cat_combo.get(), opis)
        self.destroy()
        if self.on_save_callback:
            self.on_save_callback()

class ExcelImportDialog(tk.Toplevel):
    TARGET_FIELDS = [  # noqa: RUF012
        ("model", "Model telefonu *"),
        ("nr_tel", "Numer telefonu *"),
        ("nr_sim", "Numer SIM"),
        ("imei", "IMEI"),
        ("nr_seryjny", "Numer seryjny"),
        ("rodzaj", "Rodzaj"),
        ("osoba_uzytkujaca", "Osoba użytkująca"),
        ("osoba_odpowiedzialna", "Osoba odpowiedzialna"),
    ]

    def __init__(self, parent, file_path, on_success_callback):
        super().__init__(parent)
        self.file_path = file_path
        self.on_success_callback = on_success_callback

        self.title("Import danych z Excela")
        self.geometry("600x400")
        self.transient(parent)
        self.grab_set()

        self.wb = None
        self.sheet = None
        self.header = []
        self.combos = {}

        if self.load_excel_headers():
            self.build_ui()

    def load_excel_headers(self):
        try:
            self.wb = openpyxl.load_workbook(self.file_path, data_only=True)
            self.sheet = self.wb.active

            first_row = next(self.sheet.iter_rows(values_only=True), None)
            if not first_row:
                messagebox.showerror("Błąd", "Plik Excel jest pusty.")
                self.destroy()
                return False

            self.header = [str(val).strip() if val is not None else f"Kolumna {idx+1}" for idx, val in enumerate(first_row)]
            
            return True
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Błąd", f"Nie można wczytać pliku Excel: {e}")
            self.destroy()
            return False

    def build_ui(self):
        ttk.Label(
            self,
            text="Przypisz kolumny z arkusza Excel do odpowiednich pól.\nPola oznaczone gwiazdką (*) są wymagane.",
            padding=10,
            justify=tk.LEFT
        ).pack(anchor=tk.W)

        form_frame = ttk.Frame(self, padding=(15, 5))
        form_frame.pack(fill=tk.BOTH, expand=True)

        options = ["[Ignoruj / brak]"] + self.headers

        for idx, (field_key, field_label) in enumerate(self.TARGET_FIELDS):
            ttk.Label(form_frame, text=field_label).grid(row=idx, column=0, sticky=tk.W, pady=4, padx=5)
            
            combo = ttk.Combobox(form_frame, values=options, state="readonly", width=28)
            
            # Próba automatycznego dopasowania po nazwie
            matched = False
            for header in self.headers:
                clean_field = field_label.lower().replace("*", "").strip()
                if header.lower() in clean_field or clean_field in header.lower():
                    combo.set(header)
                    matched = True
                    break
            if not matched:
                combo.set("[Ignoruj / brak]")

            combo.grid(row=idx, column=1, sticky=tk.EW, pady=4, padx=5)
            self.combos[field_key] = combo

        form_frame.columnconfigure(1, weight=1)

        btn_box = ttk.Frame(self, padding=10)
        btn_box.pack(fill=tk.X)
        ttk.Button(btn_box, text="Anuluj", command=self.destroy).pack(side=tk.RIGHT, padx=5)
        ttk.Button(btn_box, text="Importuj dane", command=self.import_data).pack(side=tk.RIGHT, padx=5)

    def import_data(self):
        mapping = {}
        for key, combo in self.combos.items():
            selected = combo.get()
            if selected != "[Ignoruj / brak]":
                mapping[key] = self.headers.index(selected)

        if "model" not in mapping or "nr_tel" not in mapping:
            messagebox.showwarning(
                "Brakujące mapowanie",
                "Wymagane jest przypisanie kolumn: 'Model telefonu' oraz 'Numer telefonu'.",
                parent=self
            )
            return

        rows = list(self.sheet.iter_rows(values_only=True))
        data_rows = rows[1:]

        phone_records = []
        for row in data_rows:
            if not any(row):
                continue

            record = {}
            for field_key, col_idx in mapping.items():
                val = row[col_idx] if col_idx < len(row) else ""
                record[field_key] = str(val) if val is not None else ""
            phone_records.append(record)

        if not phone_records:
            messagebox.showinfo("Brak danych", "Nie znaleziono żadnych danych do zaimportowania.", parent=self)
            return

        imported_count = db.bulk_insert_phones(phone_records)
        messagebox.showinfo("Sukces", f"Pomyślnie zaimportowano {imported_count} telefon(ów).", parent=self)

        self.destroy()
        if self.on_success_callback:
            self.on_success_callback()