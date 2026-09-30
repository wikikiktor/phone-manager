import os
import re
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def get_db_path():
    appdata = os.getenv("APPDATA")
    if appdata:
        app_dir = Path(appdata) / "BazaTelefonow"
    else:
        app_dir = Path.home() / ".baza_telefonow"

    app_dir.mkdir(parents=True, exist_ok=True)
    target_db = app_dir / "telefony.db"

    return str(target_db)

DB_NAME = get_db_path()

def get_connection():
    return sqlite3.connect(DB_NAME)

def format_phone_number(val):
    if not val:
        return ""
    val_str = str(val).strip()
    digits = re.sub(r"\D", "", val_str)

    if len(digits) == 9:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    elif len(digits) == 11 and digits.startswith("48"):
        d = digits[2:]
        return f"+48 {d[:3]}-{d[3:6]}-{d[6:9]}"
    else:
        return val

def init_db():

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS telefony (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            model TEXT NOT NULL,
            nr_tel TEXT NOT NULL,
            nr_sim TEXT,
            imei TEXT,
            nr_seryjny TEXT,
            rodzaj TEXT,
            osoba_uzytkujaca TEXT,
            osoba_odpowiedzialna TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS historia (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telefon_id INTEGER NOT NULL,
            data TEXT NOT NULL,
            kategoria TEXT NOT NULL,
            opis TEXT,
            FOREIGN KEY (telefon_id) REFERENCES telefony (id)
        )''')

    try:
        cursor.execute("ALTER TABLE telefony ADD COLUMN czy_usuniety INTEGER DEFAULT 0")
    except sqlite3.OperationalError:
        pass

    cursor.execute("SELECT COUNT(*) FROM telefony")
    if cursor.fetchone()[0] == 0:
        cursor.execute("""
            INSERT INTO telefony (model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
            (
                "Przykładowy Model",
                format_phone_number("123456789"),
                "SIM123456",
                "IMEI123456789",
                "SERIAL123456",
                "Rodzaj1",
                "Użytkownik1",
                "Odpowiedzialny1",
                ),)
        tel_id = cursor.lastrowid
        cursor.execute('''
            INSERT INTO historia (telefon_id, data, kategoria, opis)
            VALUES (?, ?, ?, ?)
        ''', (tel_id, 
              datetime.now().strftime("%Y-%m-%d %H:%M"),  # noqa: DTZ005
              "Dodanie", 
              "Telefon dodany do bazy danych",
              ),)

    conn.commit()
    conn.close()

def search_phones(query_str, show_deleted=False):
    search = f"%{query_str}%"
    clean_digits = re.sub(r"\D", "", query_str)
    digits_search = f"%{clean_digits}%" if clean_digits else search
    deleted_flag = 1 if show_deleted else 0

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT id, nr_tel, osoba_uzytkujaca, model FROM telefony
            WHERE (
                nr_tel LIKE ?
                OR REPLACE(nr_tel, '-', '') LIKE ? 
                OR osoba_uzytkujaca LIKE ? 
                OR model LIKE ? 
                OR imei LIKE ?
            )
            AND czy_usuniety = ?
            ORDER BY id DESC
            """,
            (search, digits_search, search, search, search, deleted_flag)
        )
        return cursor.fetchall()

def get_phone_by_id(phone_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna
            FROM telefony
            WHERE id = ?
            """,
            (phone_id,),
        )
        return cursor.fetchone()

def get_phone_history(phone_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT data, kategoria, opis
            FROM historia
            WHERE telefon_id = ?
            ORDER BY id DESC
            """,
            (phone_id,),
        )
        return cursor.fetchall()

def insert_phone(data):
    formatted_nr = format_phone_number(data.get("nr_tel",""))

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO telefony (model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                data["model"],
                formatted_nr,
                data["nr_sim"],
                data["imei"],
                data["nr_seryjny"],
                data["rodzaj"],
                data["osoba_uzytkujaca"],
                data["osoba_odpowiedzialna"],
            ),
        )
        phone_id = cursor.lastrowid
        cursor.execute(
            """
            INSERT INTO historia (telefon_id, data, kategoria, opis)
            VALUES (?, ?, ?, ?)
            """,
            (
                phone_id,
                datetime.now().strftime("%Y-%m-%d %H:%M"),  # noqa: DTZ005
                "Dodanie do bazy",
                "Zarejestrowano urządzenie w systemie.",
            ),
        )
        return phone_id

def update_phone(phone_id, data):
    data["nr_tel"]= format_phone_number(data.get("nr_tel", ""))
    
    labels = {
        "model": "Model",
        "nr_tel": "Nr telefonu",
        "nr_sim": "Nr SIM",
        "imei": "IMEI",
        "nr_seryjny": "Nr seryjny",
        "rodzaj": "Rodzaj",
        "osoba_uzytkujaca": "Osoba użytkująca",
        "osoba_odpowiedzialna": "Osoba odpowiedzialna"
    }

    category_map = {
        "osoba_uzytkujaca": "Zmiana użytkownika",
        "osoba_odpowiedzialna": "Zmiana odpowiedzialnego",
        "nr_sim": "Wymiana karty SIM",
        "rodzaj": "Zmiana statusu/rodzaju",
    }

    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute(
            """
            Select model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna
            FROM telefony
            WHERE id = ?
            """,
            (phone_id,)
        )
        old_row = cursor.fetchone()

        cursor.execute(
            """
            UPDATE telefony
            SET model = ?, nr_tel = ?, nr_sim = ?, imei = ?, nr_seryjny = ?, rodzaj = ?, osoba_uzytkujaca = ?, osoba_odpowiedzialna = ?
            WHERE id = ?
            """,
            (
                data["model"],
                data["nr_tel"],
                data["nr_sim"],
                data["imei"],
                data["nr_seryjny"],
                data["rodzaj"],
                data["osoba_uzytkujaca"],
                data["osoba_odpowiedzialna"],
                phone_id,
            ),
        )

        if old_row:
            keys = [
                "model", "nr_tel", "nr_sim", "imei",
                "nr_seryjny", "rodzaj", "osoba_uzytkujaca", "osoba_odpowiedzialna"
            ]
            now_str = datetime.now().strftime("%Y-%m-%d %H:%M")  # noqa: DTZ005
            
            for i, key in enumerate(keys):
                old_val = (old_row[i] or "").strip()
                new_val = (data[key] or "").strip()
                
                if old_val != new_val:
                    old = old_val if old_val else "[puste]"
                    new = new_val if new_val else "[puste]"
                    kategoria = category_map.get(key, "Edycja danych")
                    opis = f"Zmieniono {labels[key]}: '{old}' ➔ '{new}'"

                    cursor.execute(
                        """
                        INSERT INTO historia (telefon_id, data, kategoria, opis)
                        VALUES (?, ?, ?, ?)
                        """,
                        (phone_id, now_str, kategoria, opis),
                    )

def soft_delete_phone(phone_id):
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE telefony SET czy_usuniety = 1 WHERE id = ?", (phone_id,))
        cursor.execute(
            """
            INSERT INTO historia (telefon_id, data, kategoria, opis)
            VALUES (?, ?, ?, ?)
            """,
            (phone_id, now_str, "Kosz / Usunięcie", "Telefon przeniesiono do kosza (usunięto z aktywnej listy)."),
        )

def restore_phone(phone_id):
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE telefony SET czy_usuniety = 0 WHERE id = ?", (phone_id,))
            cursor.execute(
                """
                INSERT INTO historia (telefon_id, data, kategoria, opis)
                VALUES (?, ?, ?, ?)
                """,
                (phone_id, now_str, "Przywrócenie", "Przywrócono urządzenie z kosza do aktywnych."),
            )

def hard_delete_phone(phone_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "DELETE FROM historia WHERE telefon_id = ?", 
            (phone_id,)
        )
        cursor.execute(
            "DELETE FROM telefony WHERE id = ?", 
            (phone_id,)
        )

def add_history_entry(phone_id, category, description):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO historia (telefon_id, data, kategoria, opis)
            VALUES (?, ?, ?, ?)
            """,
            (
                phone_id,
                datetime.now().strftime("%Y-%m-%d %H:%M"),  # noqa: DTZ005
                category,
                description,
            ),
        )

def bulk_insert_phones(phone_records, history_records=None):

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")  # noqa: DTZ005
    inserted_phones_count = 0
    inserted_history_count = 0
    phone_id_map = {}

    with get_connection() as conn:
        cursor = conn.cursor()

        for data in phone_records:
            model = (data.get("model") or "").strip()
            nr_tel = format_phone_number((data.get("nr_tel") or "").strip())

            if not model or not nr_tel:
                continue

            cursor.execute(
                """
                INSERT INTO telefony (model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    model,
                    nr_tel,
                    (data.get("nr_sim") or "").strip(),
                    (data.get("imei") or "").strip(),
                    (data.get("nr_seryjny") or "").strip(),
                    (data.get("rodzaj") or "").strip(),
                    (data.get("osoba_uzytkujaca") or "").strip(),
                    (data.get("osoba_odpowiedzialna") or "").strip(),
                ),
            )
            phone_id = cursor.lastrowid
            phone_id_map[nr_tel] = phone_id
            inserted_phones_count += 1

            if not history_records:
                cursor.execute(
                    """
                    INSERT INTO historia (telefon_id, data, kategoria, opis)
                    VALUES (?, ?, ?, ?)
                    """,
                    (
                        phone_id,
                        now_str,
                        "Dodanie do bazy",
                        "Zarejestrowano urządzenie w systemie.",
                    ),
                )
        if history_records:
            for h in history_records:
                h_nr = format_phone_number(h.get("nr_tel", ""))
                phone_id = phone_id_map.get(h_nr)

                if not phone_id and h_nr:
                    cursor.execute("SELECT id FROM telefony WHERE nr_tel = ? LIMIT 1", (h_nr,))
                    row = cursor.fetchone()
                    if row:
                        phone_id = row[0]

                if phone_id:
                    cursor.execute(
                        """
                        INSERT INTO historia (telefon_id, data, kategoria, opis)
                        VALUES (?, ?, ?, ?)
                        """,
                        (
                            phone_id,
                            h.get("data") or now_str,
                            h.get("kategoria") or "Import",
                            h.get("opis") or "",
                        ),
                    )
                    inserted_history_count += 1

    return inserted_phones_count, inserted_history_count

def export_to_excel(file_path, phone_ids=None):
    wb = openpyxl.Workbook()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Elegancki granat
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    ws_phones = wb.active
    ws_phones.title = "Telefony"

    headers_phones = [
        "ID", "Model telefonu", "Numer telefonu", "Numer SIM",
        "IMEI", "Numer seryjny", "Rodzaj", "Osoba użytkująca", "Osoba odpowiedzialna"
    ]
    ws_phones.append(headers_phones)

    with get_connection() as conn:
        cursor = conn.cursor()

        if phone_ids is not None:
            if not phone_ids:
                return 0
            placeholders = ",".join("?" for _ in phone_ids)
            cursor.execute(
                f"""
                SELECT id, model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna
                FROM telefony
                WHERE id IN ({placeholders})
                ORDER BY id ASC
                """,
                phone_ids,
            )
        else:
            cursor.execute(
                """
                SELECT id, model, nr_tel, nr_sim, imei, nr_seryjny, rodzaj, osoba_uzytkujaca, osoba_odpowiedzialna
                FROM telefony
                ORDER BY id ASC
                """
            )
        phone_rows = cursor.fetchall()
    for row in phone_rows:
        ws_phones.append(list(row))

    for col_idx, cell in enumerate(ws_phones[1], start=1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for row in ws_phones.iter_rows(min_row=2, max_row=ws_phones.max_row, min_col=1, max_col=len(headers_phones)):
        for cell in row:
            cell.border = thin_border
            # ID, nr_tel, SIM, IMEI wyśrodkowane, reszta do lewej
            if cell.column in (1, 3, 4, 5):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    ws_phones.freeze_panes = "A2"
    ws_phones.auto_filter.ref = ws_phones.dimensions

    ws_hist = wb.create_sheet(title="Historia zdarzeń")
    headers_hist = ["Data", "Model telefonu", "Nr telefonu", "Kategoria", "Opis zdarzenia / uwagi"]
    ws_hist.append(headers_hist)

    with get_connection() as conn:
        cursor = conn.cursor()
        if phone_ids is not None:
            placeholders = ",".join("?" for _ in phone_ids)
            cursor.execute(
                f"""
                SELECT h.data, t.model, t.nr_tel, h.kategoria, h.opis
                FROM historia h
                LEFT JOIN telefony t ON h.telefon_id = t.id
                WHERE h.telefon_id IN ({placeholders})
                ORDER BY h.id DESC
                """,
                phone_ids,
            )
        else:
            cursor.execute(
                """
                SELECT h.data, t.model, t.nr_tel, h.kategoria, h.opis
                FROM historia h
                LEFT JOIN telefony t ON h.telefon_id = t.id
                ORDER BY h.id DESC
                """
            )
        hist_rows = cursor.fetchall()

    for row in hist_rows:
        ws_hist.append(list(row))

    # Stylowanie arkusza historii
    for cell in ws_hist[1]:
        cell.font = header_font
        cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        cell.alignment = center_align

    for row in ws_hist.iter_rows(min_row=2, max_row=ws_hist.max_row, min_col=1, max_col=len(headers_hist)):
        for cell in row:
            cell.border = thin_border
            if cell.column in (1, 3, 4):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    ws_hist.freeze_panes = "A2"
    ws_hist.auto_filter.ref = ws_hist.dimensions

    for ws in (ws_phones, ws_hist):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)
    return len(phone_rows)

def get_phones_count(show_deleted=False):
    deleted_flag = 1 if show_deleted else 0
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM telefony WHERE czy_usuniety = ?", (deleted_flag,))
        return cursor.fetchone()[0]
