import os
from datetime import datetime

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:
    print("Moduł 'python-docx' nie jest zainstalowany. Zainstaluj go, aby korzystać z funkcji eksportu do Worda.")
    Document = None

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
except ImportError:
    print("Moduł 'reportlab' nie jest zainstalowany. Zainstaluj go, aby korzystać z funkcji eksportu do PDF.")
    SimpleDocTemplate = None

def _get_pdf_font():
    fonts = [
        ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("Calibri", "C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf"),
        ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for name, reg, bold in fonts:
        if os.path.exists(reg):
            try:
                pdfmetrics.registerFont(TTFont(name, reg))
                b_name = f"{name}-Bold"
                if os.path.exists(bold):
                    pdfmetrics.registerFont(TTFont(b_name, bold))
                else:
                    pdfmetrics.registerFont(TTFont(b_name, reg))
                return name, b_name
            except Exception:
                pass
    return "Helvetica", "Helvetica-Bold"

def _prepare_data(phone_data, employer_data, protocol_type):
    model = phone_data[0] or ""
    nr_tel = phone_data[1] or ""
    nr_sim = phone_data[2] or ""
    imei = phone_data[3] or ""
    nr_seryjny = phone_data[4] or ""
    osoba = phone_data[6] or ""
    wyposazenie = (phone_data[9] or "").lower()
    stan_baterii = (phone_data[10] or "").strip().lower()

    nazwa_firmy = employer_data[0] if employer_data and employer_data[0] else ".........................................................."
    adres_firmy = employer_data[1] if employer_data and employer_data[1] else "......................................................................."
    nip_firmy = employer_data[2] if employer_data and employer_data[2] else "..........................................................................."

    today = datetime.now().strftime("%d.%m.%Y")

    has_ladowarka = "ładowarka" in wyposazenie or "ladowarka" in wyposazenie
    has_kabel = "kabel" in wyposazenie or "usb" in wyposazenie
    has_etui = "etui" in wyposazenie
    has_szklo = "szkło" in wyposazenie or "szyklo" in wyposazenie

    is_dobry = "dobry" in stan_baterii and "brak" not in stan_baterii
    is_zadowalajacy = "zadowalający" in stan_baterii or "zadowalajacy" in stan_baterii
    is_slaby = "słaby" in stan_baterii or "slaby" in stan_baterii

    return {
        "type": protocol_type,
        "date": today,
        "nazwa_firmy": nazwa_firmy,
        "adres_firmy": adres_firmy,
        "nip_firmy": nip_firmy,
        "osoba": osoba or "...........................................................................",
        "model": model or "......................................................",
        "imei": imei or "........................................................................",
        "nr_seryjny": nr_seryjny or "....................................................................",
        "nr_tel_sim": f"{nr_tel} (SIM: {nr_sim})" if nr_sim else (nr_tel or "........................................................"),
        "has_ladowarka": has_ladowarka,
        "has_kabel": has_kabel,
        "has_etui": has_etui,
        "has_szklo": has_szklo,
        "is_dobry": is_dobry,
        "is_zadowalajacy": is_zadowalajacy,
        "is_slaby": is_slaby,
    }

def generate_docx(file_path, phone_data, employer_data, protocol_type):
    if Document is None:
        raise ImportError("Moduł 'python-docx' nie jest zainstalowany.")

    d = _prepare_data(phone_data, employer_data, protocol_type)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    def p(text="", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3, line_spacing=0.9, font_size=10):
        par = doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = line_spacing
        if text:
            run = par.add_run(text)
            run.bold = bold
            run.font.name = "Palatino Linotype"
            run.font.size = Pt(font_size)
        return par

    c_box = lambda checked: "☑" if checked else "☐"

    is_przekazanie = d["type"] == "przekazanie"
    tytul = "PROTOKÓŁ PRZEKAZANIA TELEFONU SŁUŻBOWEGO" if is_przekazanie else "PROTOKÓŁ ZWROTU TELEFONU SŁUŻBOWEGO"
    p(tytul, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(f"sporządzony w dniu {d['date']}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    #1. Dane pracodawcy
    p("1. Dane pracodawcy:", bold=True, space_after=2)
    p(f"Nazwa firmy: {d['nazwa_firmy']}")     
    p(f"Adres: {d['adres_firmy']}")
    p(f"NIP: {d['nip_firmy']}", space_after=6)

    #2. Dane pracownika
    p("2. Dane pracownika", bold=True, space_after=2)
    p(f"Imię i nazwisko: {d['osoba']}", space_after=6)

    #3. Przedmiot przekazania
    sekcja_3 = "3. Przedmiot przekazania" if is_przekazanie else "3. Przedmiot zwrotu"
    p(sekcja_3, bold=True, space_after=2)
    opis_przedmiotu = (
        "Pracodawca przekazuje pracownikowi do użytkowania następujące mienie służbowe:"
        if is_przekazanie else
        "Pracownik zwraca pracodawcy następujące mienie służbowe:"
    )
    p(opis_przedmiotu, space_after=2)
    p(f"Marka i model telefonu: {d['model']}")
    p(f"Numer IMEI: {d['imei']}")
    p(f"Numer seryjny: {d['nr_seryjny']}")
    p(f"Numer telefonu (SIM): {d['nr_tel_sim']}", space_after=6)

    #4. Wyposażenie dodatkowe
    p("4. Wyposażenie dodatkowe", bold=True, space_after=2)
    p(f"Ładowarka {c_box(d['has_ladowarka'])}")
    p(f"kabel USB {c_box(d['has_kabel'])}")
    p(f"etui {c_box(d['has_etui'])}")
    p(f"szkło ochronne {c_box(d['has_szklo'])}", space_after=4)

    stan_tekst = "Stan techniczny urządzenia w chwili przekazania:" if is_przekazanie else "Stan techniczny urządzenia w chwili zwrotu:"
    p(f"{stan_tekst} ............................................................................................................................................................................................................................................................................................................................................................................................................", space_after=4, bold=True)

    bat_str = f"Stan baterii: Dobry {c_box(d['is_dobry'])} / Zadowalający {c_box(d['is_zadowalajacy'])} / Słaby {c_box(d['is_slaby'])}"
    p(bat_str, space_after=8, bold=True)

    #5. Oświadczenie pracownika
    if is_przekazanie:
        p("5. Oświadczenie pracownika", bold=True, space_after=2)
        p("Ja, niżej podpisany(-a), potwierdzam odbiór wyżej opisanego telefonu służbowego wraz z wyposażeniem i oświadczam, że:")
        punkty = [
            "Otrzymałem(-am) telefon w stanie sprawnym i bez zastrzeżeń (lub zgodnie z opisem stanu technicznego wskazanym powyżej).",
            "Zobowiązuję się do użytkowania telefonu zgodnie z jego przeznaczeniem.",
            "Zobowiązuję się do należytego zabezpieczenia telefonu przed utratą, zniszczeniem lub uszkodzeniem.",
            "Zobowiązuję się do niezwłocznego zgłaszania utraty, kradzieży, uszkodzenia lub awarii urządzenia.",
            "Zobowiązuję się do zwrotu telefonu wraz z całym przekazanym wyposażeniem niezwłocznie po ustaniu potrzeby jego używania lub najpóźniej w dniu zakończenia zatrudnienia, chyba że Pracodawca wskaże inny termin.",
            "Przyjmuję do wiadomości, że telefon został mi powierzony z obowiązkiem zwrotu i ponoszę odpowiedzialność materialną za powierzone mienie na zasadach określonych w kodeksie pracy.",
            "W przypadku uszkodzenia lub zniszczenia telefonu z mojej winy zobowiązuję się do pokrycia kosztów jego naprawy lub, jeśli naprawa nie będzie możliwa, kosztów zakupu telefonu.",
        ]
        for idx, pkt in enumerate(punkty, start=1):
            par = p(f"{idx}. {pkt}", space_after=0.5, line_spacing=1.0, font_size=9)
            par.paragraph_format.left_indent = Inches(0.15)
        p("", space_after=12)

        # Podpisy
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "............................................................"
        table.rows[0].cells[1].text = "……………………………………."
        table.rows[1].cells[0].text = "Przekazujący (Pracodawca):"
        table.rows[1].cells[1].text = "Odbierający (Pracownik):"
    else:
        p("5. Potwierdzenie odbioru przez pracodawcę", bold=True, space_after=2)
        p("Pracodawca potwierdza odbiór telefonu służbowego wraz z wyposażeniem wskazanym w niniejszym protokole. Strony potwierdzają, że stan zwracanego mienia oraz ewentualne braki, uszkodzenia lub inne uwagi zostały opisane powyżej.", space_after=20)

        # Podpisy
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "............................................................"
        table.rows[0].cells[1].text = "……………………………………."
        table.rows[1].cells[0].text = "Odbierający / Pracodawca"
        table.rows[1].cells[1].text = "Zwracający / Pracownik"

    doc.save(file_path)

def generate_pdf(file_path, phone_data, employer_data, protocol_type):
    if SimpleDocTemplate is None:
        raise ImportError("Moduł 'reportlab' nie jest zainstalowany.")

    d = _prepare_data(phone_data, employer_data, protocol_type)
    font_reg, font_bold = _get_pdf_font()

    doc = SimpleDocTemplate(
        file_path, 
        pagesize=A4, 
        rightMargin=40, 
        leftMargin=40, 
        topMargin=35, 
        bottomMargin=35,
    )

    style_title = ParagraphStyle("Tytul", fontName=font_bold, fontSize=12, alignment=1, spaceAfter=2)
    style_sub = ParagraphStyle("Sub", fontName=font_bold, fontSize=9, alignment=1, spaceAfter=8)
    style_sec = ParagraphStyle("Sec", fontName=font_bold, fontSize=9.5, spaceBefore=12, spaceAfter=2)
    style_txt = ParagraphStyle("Txt", fontName=font_reg, fontSize=8.5, leading=11, spaceAfter=2)
    style_txt_b = ParagraphStyle("TxtB", fontName=font_bold, fontSize=8.5, leading=11, spaceAfter=2)
    style_bullet = ParagraphStyle("Bullet", fontName=font_reg, fontSize=8, leading=10, leftIndent=10, spaceAfter=1.5)

    c_box = lambda checked: "[X]" if checked else "[ ]"
    is_przekazanie = d["type"] == "przekazanie"
    story = []

    tytul = "PROTOKÓŁ PRZEKAZANIA TELEFONU SŁUŻBOWEGO" if is_przekazanie else "PROTOKÓŁ ZWROTU TELEFONU SŁUŻBOWEGO"
    story.append(Paragraph(tytul, style_title))
    story.append(Paragraph(f"sporządzony w dniu {d['date']}", style_sub))

    # 1. Dane pracodawcy
    story.append(Paragraph("1. Dane pracodawcy", style_sec))
    story.append(Paragraph(f"Nazwa firmy: {d['nazwa_firmy']}", style_txt))
    story.append(Paragraph(f"Adres: {d['adres_firmy']}", style_txt))
    story.append(Paragraph(f"NIP: {d['nip_firmy']}", style_txt))

    # 2. Dane pracownika
    story.append(Paragraph("2. Dane pracownika", style_sec))
    story.append(Paragraph(f"Imię i nazwisko: {d['osoba']}", style_txt))

    # 3. Przedmiot
    sekcja_3 = "3. Przedmiot przekazania" if is_przekazanie else "3. Przedmiot zwrotu"
    story.append(Paragraph(sekcja_3, style_sec))
    opis_przedmiotu = (
        "Pracodawca przekazuje pracownikowi do użytkowania następujące mienie służbowe:"
        if is_przekazanie else
        "Pracownik zwraca pracodawcy następujące mienie służbowe:"
    )
    story.append(Paragraph(opis_przedmiotu, style_txt))
    story.append(Paragraph(f"Marka i model telefonu: {d['model']}", style_txt))
    story.append(Paragraph(f"Numer IMEI: {d['imei']}", style_txt))
    story.append(Paragraph(f"Numer seryjny: {d['nr_seryjny']}", style_txt))
    story.append(Paragraph(f"Numer telefonu (SIM): {d['nr_tel_sim']}", style_txt))

    # 4. Wyposażenie
    story.append(Paragraph("4. Wyposażenie dodatkowe", style_sec))
    story.append(Paragraph(f"Ładowarka {c_box(d['has_ladowarka'])}", style_txt))
    story.append(Paragraph(f"kabel USB {c_box(d['has_kabel'])}", style_txt))
    story.append(Paragraph(f"etui {c_box(d['has_etui'])}", style_txt))
    story.append(Paragraph(f"szkło ochronne {c_box(d['has_szklo'])}", style_txt))

    stan_tekst = "Stan techniczny urządzenia w chwili przekazania:" if is_przekazanie else "Stan techniczny urządzenia w chwili zwrotu:"
    story.append(Paragraph(f"{stan_tekst}", style_txt_b))
    story.append(Paragraph("..........................................................................................................................................................................................................................................................................................................................................................................................................................................", style_txt))

    bat_str = f"Stan baterii: Dobry {c_box(d['is_dobry'])} / Zadowalający {c_box(d['is_zadowalajacy'])} / Słaby {c_box(d['is_slaby'])}"
    story.append(Paragraph(bat_str, style_txt_b))

    # 5. Oświadczenie / Potwierdzenie
    if is_przekazanie:
        story.append(Paragraph("5. Oświadczenie pracownika", style_sec))
        story.append(Paragraph("Ja, niżej podpisany(-a), potwierdzam odbiór wyżej opisanego telefonu służbowego wraz z wyposażeniem i oświadczam, że:", style_txt))
        punkty = [
            "Otrzymałem(-am) telefon w stanie sprawnym i bez zastrzeżeń (lub zgodnie z opisem stanu technicznego wskazanym powyżej).",
            "Zobowiązuję się do użytkowania telefonu zgodnie z jego przeznaczeniem.",
            "Zobowiązuję się do należytego zabezpieczenia telefonu przed utratą, zniszczeniem lub uszkodzeniem.",
            "Zobowiązuję się do niezwłocznego zgłaszania utraty, kradzieży, uszkodzenia lub awarii urządzenia.",
            "Zobowiązuję się do zwrotu telefonu wraz z całym przekazanym wyposażeniem niezwłocznie po ustaniu potrzeby jego używania lub najpóźniej w dniu zakończenia zatrudnienia, chyba że Pracodawca wskaże inny termin.",
            "Przyjmuję do wiadomości, że telefon został mi powierzony z obowiązkiem zwrotu i ponoszę odpowiedzialność materialną za powierzone mienie na zasadach określonych w kodeksie pracy.",
            "W przypadku uszkodzenia lub zniszczenia telefonu z mojej winy zobowiązuję się do pokrycia kosztów jego naprawy lub, jeśli naprawa nie będzie możliwa, kosztów zakupu telefonu.",
        ]
        for idx, pkt in enumerate(punkty, start=1):
            story.append(Paragraph(f"{idx}. {pkt}", style_bullet))
        story.append(Spacer(1, 15))

        sig_data = [
            [Paragraph("............................................................", style_txt), Paragraph("…………………………………….", style_txt)],
            [Paragraph("Przekazujący (Pracodawca):", style_txt), Paragraph("Odbierający (Pracownik):", style_txt)],
        ]
    else:
        story.append(Paragraph("5. Potwierdzenie odbioru przez pracodawcę", style_sec))
        story.append(Paragraph("Pracodawca potwierdza odbiór telefonu służbowego wraz z wyposażeniem wskazanym w niniejszym protokole. Strony potwierdzają, że stan zwracanego mienia oraz ewentualne braki, uszkodzenia lub inne uwagi zostały opisane powyżej.", style_txt))
        story.append(Spacer(1, 25))

        sig_data = [
            [Paragraph("............................................................", style_txt), Paragraph("…………………………………….", style_txt)],
            [Paragraph("Odbierający / Pracodawca", style_txt), Paragraph("Zwracający / Pracownik", style_txt)],
        ]

    sig_table = Table(sig_data, colWidths=[260, 250])
    sig_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(sig_table)

    doc.build(story)